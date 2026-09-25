"""Builds parser test fixtures from the real demo documents (no invented
content): a DOCX of the 2026 regulations using Word heading styles and a
Word table, and an image-only (scanned) PDF of the regulations 2025 text.

    python tests/make_fixtures.py
"""
import re
from pathlib import Path

import docx
import fitz

HERE = Path(__file__).parent / "fixtures"
DEMO = Path(__file__).resolve().parents[2] / "demo-data"


def make_docx():
    text = (DEMO / "university_regulations_2026.txt").read_text(encoding="utf-8")
    paras = [p.strip() for p in re.split(r"\n\s*\n", text) if p.strip()]
    d = docx.Document()
    d.add_heading(paras[0], level=0)  # Title style
    for p in paras[1:]:
        if "\n" not in p and len(p) < 60:
            d.add_heading(p, level=1)
        else:
            d.add_paragraph(" ".join(p.split()))
    # the grading weights from the same document, as a Word table
    d.add_heading("Grading Weights (table form)", level=2)
    t = d.add_table(rows=4, cols=2)
    for row, (a, b) in zip(t.rows, [("Component", "Weight"), ("Continuous assessment", "25%"),
                                     ("Mid-semester examination", "25%"), ("End-semester examination", "50%")]):
        row.cells[0].text, row.cells[1].text = a, b
    d.save(HERE / "university_regulations_2026.docx")


def make_scanned_pdf():
    """Renders the text to a page, rasterizes it, and saves a PDF that holds
    only the image — no text layer, like a scanner produces."""
    text = (DEMO / "university_regulations_2025.txt").read_text(encoding="utf-8")
    src = fitz.open()
    page = src.new_page()
    page.insert_textbox(fitz.Rect(56, 56, 540, 800), text, fontsize=10, fontname="helv")
    pix = page.get_pixmap(dpi=200)
    out = fitz.open()
    p = out.new_page(width=page.rect.width, height=page.rect.height)
    p.insert_image(p.rect, stream=pix.tobytes("png"))
    out.save(HERE / "scanned_regulations_2025.pdf")


if __name__ == "__main__":
    HERE.mkdir(exist_ok=True)
    make_docx()
    make_scanned_pdf()
    print("fixtures written to", HERE)
